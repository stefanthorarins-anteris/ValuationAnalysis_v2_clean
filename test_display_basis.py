# -*- coding: utf-8 -*-
"""WHICH PRICE, WHICH EARNINGS, AND NEVER A NUMBER WHERE THERE WAS A BLANK.

The 2026-09-03 display-basis findings, all S2 -- a right pick carrying wrong information --
all in `postBo.py`, all fixed together because they are one defect wearing four costumes:

  1.  THE SAME FACT, TWO NUMBERS.  TNK, the #1 name of the shipped 2026-08-31 run, published
      `PE-ratio = 2.4980` in `AggScoreTop100` (the newest single quarter annualised, at the
      period-end price, unlabelled) and `5.26` in the XLSX (FMP's fiscal-year figure).
      `GrahamNumberToPrice` was `2.2941` in the CSV (our Stage-2 sixteen-quarter mean) and
      `1.31` in the XLSX (FMP's annual Graham over the DCF endpoint's quote).  Three price
      bases, two earnings bases, zero labels.
  2.  A MISNAMED VENDOR COLUMN.  The XLSX's `Price-to-fair value` is FMP's `priceFairValue`,
      which is P/B under another name -- MEASURED on the shipped 2026-08-31 workbook: it is
      character-identical to the `Price-to-book` column on 120 of 120 populated cells across
      all 20 name pages, and differs on none.
  3.  A BLANK PRINTED AS A ZERO.  `km.dividendYield.fillna(0)` made an absent dividend figure
      render as `0.0000` -- an affirmative claim, made out of a gap, on the page the CEO
      reviews.  40 of the 120 populated dividend cells on that workbook print `0.0000` and
      nothing distinguishes a real zero from a filled one.
  4.  A DATE WITH THE DAY AND MONTH SWAPPED.  `strftime('%Y-%d-%m')` asked
      `v4/sector_price_earning_ratio` for `2026-05-07` when it meant `2026-07-05`, every run,
      all year -- always a real, parseable, in-range date, so nothing ever reported a failure.

THE RULING (CEO, 2026-09-03), composed from two standing rules.  (a) COMPUTE, DO NOT CONSUME:
one basis and it is OURS -- TTM EPS at the price the ratio is taken at -- and the vendor's P/E
and Graham are DROPPED FROM THE DELIVERABLES ENTIRELY rather than shown beside ours.  (b) Where
two of OUR OWN bases both apply, publish the LESS FLATTERING one: at a freight peak TTM gives
TNK the higher P/E, so that is the one we lead with.  The accepted caveat, recorded because it
is a real cost and not a footnote: less flattering is NOT always more accurate -- a company
whose newest quarter just collapsed shows a worse single-quarter P/E than TTM -- and for a
preserve-capital list that is the right error to make.  The basis is labelled on the artifact
so the CEO can overrule it by eye.

EVERY TEST HERE WAS RUN AGAINST THE UNFIXED `postBo.py` AND OBSERVED TO FAIL.  The pre-fix
file was copied aside byte-identically and loaded onto a shadow `sys.path`; no `git stash` and
no `git checkout` touched this shared working tree.  The failure text is recorded in the
hand-back.  Where a test fails pre-fix with an `AttributeError` (the API is new) that proves
NOVELTY, not behaviour, and it is said so on the test -- the decisive ones are the tests that
drive the REAL `writeBoAggToCSV` / `createPresentation` and fail on an assertion about a
published cell.

NO NETWORK.  `gdg.safe_json_list` is replaced wholesale -- the same seam
`test_post_fetch_hardening` and `test_deliverable_gate` use -- and every artifact is written to
an absolute `tmp_path`, so the repo write guard is not tripped.
"""
import io
import os
import inspect

import numpy as np
import pandas as pd
import pytest

import postBo as pb
import getData_gen as gdg


REPO = os.path.dirname(os.path.abspath(__file__))


# =========================================================================== #
#  FIXTURES -- THE SHAPE OF THE REAL ARTIFACT                                  #
# =========================================================================== #
#  A fixture a fifth of the real width let a bug through last week
#  (`test_deliverable_gate._AGGSCORE_COLUMNS` carries that lesson in its own comment), so the
#  panel below carries every column `writeBoAggToCSV` reads -- the P/E pair, the Graham pair,
#  the dividend window, the current-ratio pair and the refusal stamp -- and the numbers for
#  TNK are the REAL ones off the saved 2026-08-31 run
#  (`Boresults_dic-fmp_stock_CUR6K_all_2026-08-31...pickle`, `cdx_dftop100`, newest row):
#
#      price (marketCap/weightedAverageShsOut) = 64.87   epsTTM = 17.012959
#      earningsYield                           = 0.100081 (per quarter, rpy = 4)
#
#  so the two readings that were in circulation are both reproducible from one fixture:
#      1 / (4 x 0.100081) = 2.4980   <- what shipped
#      64.87 / 17.012959  = 3.8130   <- what ships now
_TNK_PRICE = 64.87
_TNK_EPSTTM = 17.012959
_TNK_EARNYIELD = 0.100081
_TNK_PE_OLD = 2.4980            # the shipped 2026-08-31 cell
_TNK_PE_NEW = 3.8130            # the same panel row, on the TTM basis


def _panel_rows(source, price, eps_ttm, earn_yield, n=8, refused='', ccy='USD',
                dividends=-1.0e5):
    """`n` quarterly rows for one source, newest first, on the real panel's schema."""
    rows = []
    for k in range(n):
        rows.append({
            'source': source,
            'date': pd.Timestamp('2026-04-01') - pd.DateOffset(months=3 * k),
            'period': 'Q%d' % (k % 4 + 1),
            'reportedCurrency': ccy,
            'reportingFrequency': 'quarterly',
            'price': price,
            'epsTTM': eps_ttm,
            'earningsYield': earn_yield,
            'grahamNumber': 161.667045,
            'marketCap': price * 1.0e6,
            'weightedAverageShsOut': 1.0e6,
            'netIncome': earn_yield * price * 1.0e6,
            'dividendsPaid': dividends,
            'bookValuePerShare': price * 0.9,
            'totalCurrentAssets': 2.0e8,
            'totalCurrentLiabilities': 1.0e8,
            'pbRatio': 1.1,
            'totalAssets': 4.0e8,
            'totalLiabilities': 2.0e8,
            'totalStockholdersEquity': 2.0e8,
            'revenue': 1.0e8,
            'grossProfit': 4.0e7,
            #  Section 5 stamps the RAW field it refused, on that row only.
            'sanityRefusedFields': refused if k == 0 else '',
        })
    return rows


@pytest.fixture
def panel():
    """Four names covering every outcome the published P/E has.

    `REFUSD` is the shape that matters and it is NOT invented: the price-scale rule refuses
    `marketCap` and leaves the DERIVED `price` column -- computed earlier, from that very
    market cap -- finite and in place.  See `_pe_cell`'s note on why the refusal is tested
    before the value.
    """
    rows = []
    rows += _panel_rows('TNK', _TNK_PRICE, _TNK_EPSTTM, _TNK_EARNYIELD)
    rows += _panel_rows('LOSSY', 20.0, -3.0, -0.05)
    rows += _panel_rows('REFUSD', 10.0, 2.0, 0.05, refused='|marketCap|')
    rows += _panel_rows('CLEAN2', 50.0, 5.0, 0.05)
    return pd.DataFrame(rows)


_SYMBS = ['TNK', 'LOSSY', 'REFUSD', 'CLEAN2']

#  A HEALTHY 200 CARRYING THE VENDOR'S OWN OPINION OF THE SAME QUANTITIES.  The P/E is
#  deliberately far from ours and positive, so a fallback -- if one survived -- publishes a
#  number nothing else in the run agrees with, and the test can see it.
#  THE ANNUAL LEGS CARRY SIX ROWS AND THEIR OWN `symbol`/`date`, because the XLSX page is a
#  MULTI-YEAR TABLE and a one-row fixture would not be the shape of the artifact -- exactly the
#  width lesson `test_deliverable_gate` records.  Six is what FMP returned for TNK on the
#  2026-08-31 workbook, and the vendor P/E values are its real ones.
_YEARS = ['2025-12-31', '2024-12-31', '2023-12-31', '2022-12-31', '2021-12-31', '2020-12-31']
_VENDOR_PE = [5.2631, 3.4009, 3.2832, 4.5712, -1.5202, 4.2510]     # the shipped TNK column
_VENDOR_PTB = [0.9046, 0.7794, 1.1012, 0.9789, 0.4394, 0.3441]

_VENDOR = {
    'ratios': [{'symbol': 'TNK', 'date': d, 'priceEarningsRatio': 99.9999,
                'currentRatio': 2.0, 'grossProfitMargin': 0.4,
                #  the misnamed twin: FMP's `priceFairValue` IS the P/B, value for value
                'priceFairValue': p} for d, p in zip(_YEARS, _VENDOR_PTB)],
    'profile': [{'price': 88.70, 'beta': -0.229, 'sector': 'Industrials',
                 'currency': 'USD', 'companyName': 'Teekay Tankers',
                 'industry': 'Marine Shipping', 'mktCap': 3.0e9}],
    'rating': [{'ratingRecommendation': 'Strong Buy'}],
    'key-metrics': [{'peRatio': pe, 'ptbRatio': ptb, 'currentRatio': 7.9788,
                     'earningsYield': 0.189976, 'dividendYield': 0.037325,
                     'grahamNumber': 161.667045, 'freeCashFlowPerShare': 3.0}
                    for pe, ptb in zip(_VENDOR_PE, _VENDOR_PTB)],
    'dcf': [{'dcf': 100.0, 'Stock Price': 88.7}],
    'peers': [{'peersList': ['STNG', 'FRO']}],
    'cash-flow': [{'freeCashFlow': 3.0e6} for _ in _YEARS],
    'sector-PE': [{'sector': 'Industrials', 'pe': 21.5}],
}


def _install_offline_vendor(monkeypatch, recorder=None, overrides=None):
    """Replace the ONE network seam.  `recorder` collects every URL asked for."""
    table = dict(_VENDOR)
    table.update(overrides or {})

    def fake(url, label='', *a, **kw):
        if recorder is not None:
            recorder.append(url)
        for key, payload in table.items():
            if key in str(label) or key in str(url):
                return payload
        return []

    monkeypatch.setattr(gdg, 'safe_json_list', fake)
    monkeypatch.setattr(pb.gdg, 'safe_json_list', fake)
    return table


def _write_csv(tmp_path, panel_df, symbs=None, monkeypatch=None, recorder=None,
               overrides=None):
    """Drive the REAL `writeBoAggToCSV` offline and return the artifact as TEXT.

    `dtype=str, keep_default_na=False`: the assertions are about the CHARACTERS in the file,
    not about what pandas can re-parse them into.  `3.8130` and `3.813` are the same float and
    a different cell -- and, more to the point, pandas' default NA list contains the literal
    string `'NaN'`, which is this CSV's own absence SENTINEL, so a plain read turns the
    pipeline's deliberate marker into a missing value and the test can no longer tell the two
    apart.  That is the same class of thing the fixes below are about.
    """
    symbs = symbs or _SYMBS
    _install_offline_vendor(monkeypatch, recorder, overrides)
    fb_df = pd.DataFrame({'source': symbs,
                          'AggScore': list(np.linspace(0.9, 0.5, len(symbs)))})
    raw_df = pd.DataFrame({'source': symbs,
                           'grahamNumberToPrice': [2.2941] + [0.5] * (len(symbs) - 1),
                           'CycleHeat': list(range(len(symbs)))})
    empty = pd.DataFrame(columns=['source'])
    out = str(tmp_path / 'AggScoreTop100-2026-08-31_fmp_stock_CUR6K.csv')
    pb.writeBoAggToCSV(fb_df, empty, empty, 'http://base/', 'KEY', len(symbs), out,
                       flag_df=None, raw_df=raw_df,
                       universe_stamp={'universe': 'stock_CUR6K'},
                       cdx_df=panel_df, missing_fill_df=None, run_date='2026-08-31')
    return pd.read_csv(out, dtype=str, keep_default_na=False).set_index('source')


def _write_xlsx(tmp_path, monkeypatch, recorder=None, overrides=None, symbs=('TNK',)):
    """Drive the REAL `createPresentation` offline and return the first sheet as a grid."""
    import openpyxl
    _install_offline_vendor(monkeypatch, recorder, overrides)
    fb_df = pd.DataFrame({'source': list(symbs),
                          'AggScore': list(np.linspace(0.9, 0.5, len(symbs)))})
    empty = pd.DataFrame(columns=['source'])
    fname = str(tmp_path / 'PresentationTop20-2026-08-31_fmp_stock_CUR6K.xlsx')
    rep = pb.createPresentation(fb_df, empty, empty, 'http://base/', 'KEY', len(symbs),
                                fname, 10, flag_df=None, bands=None)
    assert rep['written'] == len(symbs), rep
    wb = openpyxl.load_workbook(fname)
    ws = wb[symbs[0]]
    grid = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=4)]
    wb.close()
    return grid


# =========================================================================== #
#  1.  ONE P/E, ON THE TTM BASIS, IN THE ARTIFACT ITSELF                       #
# =========================================================================== #

def test_the_published_PE_is_TTM_EPS_over_the_PERIOD_END_price(tmp_path, monkeypatch, panel):
    """*** THE DEFECT: TNK shipped `PE-ratio = 2.4980` at RANK 1 on 2026-08-31. ***

    That is `1 / (rpy x earningsYield)` on the NEWEST SINGLE QUARTER, annualised -- a crude
    tanker owner at a freight peak, publishing one peak quarter as if it earned that four
    times a year.  Its own panel row gives a trailing-twelve-month EPS of 17.012959 against
    an annualised-quarter 25.97, so the shipped cell was a third too cheap.

    DRIVEN THROUGH THE REAL WRITER, not asserted on a helper: the defect was in what reached
    the CSV, so the test reads the CSV.  A helper-level assertion would have passed while the
    caller published something else, which is how the two-basis column survived four reviews.

    CANNOT DETECT: whether TTM is the RIGHT window for every name.  It is the LESS FLATTERING
    one at a peak, which is the rule the CEO chose, and the accepted cost is a collapsed
    newest quarter reading better here than it would on a single-quarter basis.
    """
    got = _write_csv(tmp_path, panel, monkeypatch=monkeypatch)
    assert got.loc['TNK', 'PE-ratio'] == '%.4f' % _TNK_PE_NEW
    #  and it is NOT the reading that shipped -- stated as its own assertion so a future
    #  change back to the quarterly basis fails on the number, not on a rounding.
    assert got.loc['TNK', 'PE-ratio'] != '%.4f' % _TNK_PE_OLD
    assert abs(float(got.loc['TNK', 'PE-ratio']) - _TNK_PRICE / _TNK_EPSTTM) < 5e-5


def test_the_PE_ships_its_BASIS_in_a_column_beside_it(tmp_path, monkeypatch, panel):
    """A ratio with a price in it has to say WHICH price, on the artifact, not in the log.

    It matters more here than for Graham: the `price` column two to the left is a LIVE
    PROFILE QUOTE in the trading currency (88.70 USD for TNK) while the ratio is taken at the
    PERIOD-END panel price (64.87), so a reader who divides the two gets a number that means
    nothing.  The token names both legs.
    """
    got = _write_csv(tmp_path, panel, monkeypatch=monkeypatch)
    assert 'PE-ratio_basis' in got.columns
    assert got.loc['TNK', 'PE-ratio_basis'] == pb.PE_BASIS_TTM
    #  the token must name the EARNINGS basis and the PRICE basis, not a code path
    assert 'epsTTM' in pb.PE_BASIS_TTM and 'periodEndPrice' in pb.PE_BASIS_TTM
    #  ADJACENCY IS PART OF IT: column order is what a human reads.
    cols = list(got.columns)
    assert cols.index('PE-ratio') + 1 == cols.index('PE-ratio_basis')
    #  and the price it was NOT taken at is still on the row, still labelled with ITS currency
    assert got.loc['TNK', 'price'] == '88.7000' and got.loc['TNK', 'priceCurrency'] == 'USD'


def test_the_vendor_priceEarningsRatio_NEVER_reaches_the_CSV(tmp_path, monkeypatch, panel):
    """*** The fallback the CEO ruled out on 2026-09-03. ***

    A loss-maker's panel cannot answer, and the vendor's `priceEarningsRatio` for that same
    name used to be published in its place behind a positive-sign test.  The vendor here
    answers 99.9999 -- positive, so the old sign test admits it -- and the cell must still be
    blank, because a second unlabelled basis in a column whose whole promise is that it has
    ONE is the defect, not a gap in coverage.

    WHAT IS LOST IS NAMED RATHER THAN GLOSSED: a genuine loss-maker whose vendor P/E was
    positive on a fiscal-year window now shows a blank instead of a number.
    """
    got = _write_csv(tmp_path, panel, monkeypatch=monkeypatch)
    assert got.loc['LOSSY', 'PE-ratio'] == 'NaN'
    assert got.loc['LOSSY', 'PE-ratio_basis'] == pb.PE_BASIS_NO_TTM_EPS
    #  the vendor's number must appear NOWHERE in the artifact, under any column
    body = ','.join(str(v) for v in got.to_numpy().ravel())
    assert '99.9999' not in body
    #  ... and the decision function has no vendor parameter left to pass one through
    assert 'vendor' not in inspect.signature(pb._pe_cell).parameters


def test_a_REFUSED_input_and_a_LOSS_MAKER_get_DIFFERENT_basis_tokens(tmp_path, monkeypatch,
                                                                     panel):
    """Both are blank cells; they are not the same fact about the run.

    "The pipeline judged this name's own numbers to contradict each other" and "this company
    lost money over the trailing year" are different things to know about a name in the top
    20, and a bare `NaN` says neither.  This is the `_cr_refused` idiom the `currentRatio`
    column already uses, applied to the P/E.
    """
    got = _write_csv(tmp_path, panel, monkeypatch=monkeypatch)
    assert got.loc['REFUSD', 'PE-ratio'] == 'NaN'
    assert got.loc['REFUSD', 'PE-ratio_basis'] == pb.PE_BASIS_REFUSED
    assert got.loc['LOSSY', 'PE-ratio_basis'] == pb.PE_BASIS_NO_TTM_EPS
    assert pb.PE_BASIS_REFUSED != pb.PE_BASIS_NO_TTM_EPS


def test_the_REFUSAL_beats_a_derived_value_that_OUTLIVED_it(panel):
    """*** Found while testing the new basis, and it is LIVE, not latent. ***

    `getData_fmp` stamps `price` and `epsTTM` during frame assembly
    (`stamp_frequency_and_graham`) and runs `refuse_impossible_cells` AFTER it, per ticker.
    The price-scale rule refuses `marketCap` -- not `price` -- so a refused name keeps a
    perfectly finite `price`, derived from the very market cap just judged corrupt.  Tested
    value-first (which is the order the function used until 2026-09-03) that name publishes a
    P/E built on a refused input and labels it `epsTTM/periodEndPrice`, i.e. as ours and sound.

    The fixture is exactly that shape: `REFUSD` carries `|marketCap|` on its newest row and a
    finite price and epsTTM beside it, so `_pe_ttm_ratio` DOES compute a value and the refusal
    has to be what suppresses it.
    """
    table = pb._pe_ttm_panel_table(panel)
    #  the value really is computable -- otherwise the ordering is untested
    assert pb._pe_ttm_ratio(table, 'REFUSD') == pytest.approx(5.0)
    assert pb._pe_cell(pb._pe_ttm_ratio(table, 'REFUSD'), 'REFUSD',
                       {'REFUSD'}) == ('NaN', pb.PE_BASIS_REFUSED)
    #  and the refusal is found from the panel itself, on a field the P/E is DERIVED from
    assert pb._pe_refused_sources(panel) == {'REFUSD'}
    assert 'marketCap' in pb.PE_INPUT_FIELDS and 'netIncome' in pb.PE_INPUT_FIELDS


def test_the_PE_survives_a_row_whose_VENDOR_CALLS_ALL_FAILED(tmp_path, monkeypatch, panel):
    """A degraded row must cost only what actually came from the vendor.

    The P/E now comes from the run's own panel, so a throttled `profile` / `rating` / `ratios`
    call must not blank it -- the same move `dividendYield` and `GrahamNumberToPrice` made on
    2026-08-13, and the reason `pEratioVec` left `_row_vectors`.  Every vendor call here
    returns `[]`, which is what a throttled 200 and a genuinely empty response both look like
    to this stage.
    """
    got = _write_csv(tmp_path, panel, monkeypatch=monkeypatch,
                     overrides={k: [] for k in _VENDOR})
    assert got.loc['TNK', 'PE-ratio'] == '%.4f' % _TNK_PE_NEW, (
        'a value held offline was blanked by a failed vendor call'
    )
    assert got.loc['TNK', 'PE-ratio_basis'] == pb.PE_BASIS_TTM
    #  the vendor-sourced columns on that same row DO go blank -- that is the contract
    assert got.loc['TNK', 'beta'] == 'NaN' and got.loc['TNK', 'price'] == 'NaN'


#  THE ROW-GUARD PIN IS NOT DUPLICATED HERE, DELIBERATELY.  `pEratioVec` leaving
#  `_row_vectors` (11 -> 10) is pinned by
#  `test_post_fetch_hardening.test_the_row_guard_PADS_and_so_cannot_ragged_the_vectors`,
#  which already owns that literal and carries the 12 -> 11 history beside it.  A second copy
#  of the same magic number in a second file is two things to update and one of them will be
#  missed -- which is the shape of defect this whole change is about.  That test was updated
#  in place (and now FAILS against the unfixed code, where it used to pass).


def test_the_GRAHAM_basis_tokens_name_a_BASIS_not_a_CODE_PATH(tmp_path, monkeypatch, panel):
    """`scored` / `panel-latest` told a reader where the number came from inside this program
    and nothing about which price it divides -- while the XLSX published a DIFFERENT Graham
    ratio under the same English name (2.2941 against 1.31 for TNK on 2026-08-31)."""
    for token in (pb.GRAHAM_BASIS_SCORED, pb.GRAHAM_BASIS_PANEL):
        assert 'periodEndPrice' in token, token
        assert 'ttmGraham' in token, token
    #  and the window is what distinguishes them, so the tokens must differ on it
    assert '16q-mean' in pb.GRAHAM_BASIS_SCORED
    assert pb.GRAHAM_BASIS_SCORED != pb.GRAHAM_BASIS_PANEL
    got = _write_csv(tmp_path, panel, monkeypatch=monkeypatch)
    assert got.loc['TNK', 'GrahamNumberToPrice_basis'] == pb.GRAHAM_BASIS_SCORED


# =========================================================================== #
#  2 + 3.  THE XLSX: NO VENDOR TWIN, NO MISNAMED COLUMN, NO ZERO FOR A BLANK   #
# =========================================================================== #

def test_the_XLSX_no_longer_publishes_the_vendor_PE_the_vendor_GRAHAM_or_priceFairValue(
        tmp_path, monkeypatch):
    """*** The three columns the CEO ruled out of the deliverable on 2026-09-03. ***

    `Price-to-fair value` goes for its own reason (finding 2): it is FMP's `priceFairValue`,
    which IS P/B -- character-identical to the `Price-to-book` column on 120 of 120 populated
    cells of the shipped 2026-08-31 workbook and different on none.  A vendor field that
    duplicates a computed one under a name implying a fair-value comparison is worse than a
    missing column, because the reader believes he is seeing a second, independent opinion.

    The vendor P/E and the vendor Graham go under (a) compute, do not consume.  The ONE P/E
    and the ONE Graham live in `AggScoreTop100` with their basis beside them.
    """
    grid = _write_xlsx(tmp_path, monkeypatch)
    header = _table_header(grid)
    for gone in ('PE-ratio', 'Price-to-fair value', 'Graham number to price'):
        assert gone not in header, '%r is still published on the page the CEO reviews' % gone
    #  and the vendor's own values must not appear under ANY header on the page
    body = [str(c) for row in grid[1:] for c in row if c is not None]
    assert '5.2631' not in body, "FMP's fiscal-year peRatio is still on the sheet"
    #  the columns that remain are the ones the ruling kept
    assert pb.XL_PTB_COL in header and pb.XL_PRICE_COL in header


def _table_header(grid):
    """The per-name TABLE's header row only.

    Row 1 of the sheet carries TWO things: the table's headers in columns A..N, and the
    side-block labels (`Company`, `beta`, `Market Cap`, ...) that `createPresentation` writes
    at `psdf_col = len(symb_df.columns) + 2`.  The gutter column between them is empty, so the
    table ends at the first `None` -- and taking the whole row instead is how the first
    version of this test asserted `'FMP' in 'Company'`.
    """
    out = []
    for cell in grid[0]:
        if cell is None:
            break
        out.append(cell)
    return out


def test_every_XLSX_ratio_header_NAMES_the_price_it_divides(tmp_path, monkeypatch):
    """The requirement, asserted on the rendered header row rather than on the constants.

    A reader has to be able to tell a vendor fiscal-year price from a period-end synthetic one
    from a live quote WITHOUT opening the source.  What is left in this table is FMP's
    reported ANNUAL history and nothing else, and every header says so.
    """
    grid = _write_xlsx(tmp_path, monkeypatch)
    header = _table_header(grid)
    assert header[:2] == ['Symbol', 'Date']
    for col in header[2:]:
        assert 'FMP' in col, '%r does not say whose number it is' % col
    #  the three that divide by a price say WHICH price
    for col in (pb.XL_EARNYIELD_COL, pb.XL_PTB_COL, pb.XL_DIVYIELD_COL):
        assert 'vendor price' in col, col
    assert 'DCF-endpoint quote' in pb.XL_PRICE_COL
    assert 'trading currency' in pb.XL_PRICE_COL
    #  a bare, basis-free header must not survive anywhere in the row
    assert 'Earnings yield' not in header and 'Price-to-book' not in header


def test_an_ABSENT_dividend_is_rendered_as_ABSENCE_not_as_a_ZERO(tmp_path, monkeypatch):
    """*** THE DEFECT: `km.dividendYield.fillna(0)` -> a blank prints `0.0000`. ***

    A zero is a claim -- this company pays nothing -- and it was being made out of a gap, on
    the page the CEO reviews, for a preserve-capital screen where "pays no dividend" is a real
    input to the decision.  40 of the 120 populated dividend cells on the shipped 2026-08-31
    workbook print `0.0000` and nothing on the sheet distinguishes a real zero from a filled
    one.

    The vendor here answers `null`, which is what FMP genuinely returns for plenty of non-US
    listings.  The cell must read as the sheet's own absence marker, and a REAL zero must
    still read as a zero -- an absence marker that swallows genuine zeros is the same defect
    pointing the other way.
    """
    km_null = [{'peRatio': 5.0, 'ptbRatio': 0.9, 'currentRatio': 7.9,
                'earningsYield': 0.18, 'dividendYield': None,
                'grahamNumber': 161.6, 'freeCashFlowPerShare': 3.0}] * 4
    grid = _write_xlsx(tmp_path, monkeypatch, overrides={'key-metrics': km_null})
    header = _table_header(grid)
    #  THE COLUMN IS FOUND BY WHAT IT IS, NOT BY THE NEW CONSTANT'S NAME, and that is the
    #  difference between a test of the fix and a test of the rename.  Matching
    #  `pb.XL_DIVYIELD_COL` makes this fail on the UNFIXED code with `AttributeError: no
    #  attribute XL_DIVYIELD_COL` -- which proves the constant is new and says nothing about
    #  the cell.  Matching the substring locates the column in BOTH versions, so the unfixed
    #  code fails here on the published value: `'0.0000' != '0.0000'`.
    i = next(k for k, c in enumerate(header) if 'ividend' in str(c))
    assert grid[1][i] != '0.0000', 'an absent dividend still prints as a zero'
    assert grid[1][i] == 'N/A'
    assert header[i] == pb.XL_DIVYIELD_COL

    #  ... AND A GENUINE ZERO STILL SAYS ZERO.  An absence marker that swallows real zeros is
    #  the same defect pointing the other way, and "pays no dividend" is a real input to a
    #  preserve-capital decision.
    km_zero = [dict(r, dividendYield=0.0) for r in km_null]
    grid0 = _write_xlsx(tmp_path, monkeypatch, overrides={'key-metrics': km_zero})
    assert _table_header(grid0)[i] == header[i]
    assert grid0[1][i] == '0.0000'


def test_xl_num_turns_no_absence_into_a_number():
    """The one reader, unit-tested on every shape an absence arrives in.

    PRE-FIX THIS FAILS WITH AttributeError, WHICH PROVES NOVELTY AND NOT BEHAVIOUR -- said
    plainly rather than left for a reviewer to notice.  The behavioural proof is the driven
    test above.  This one exists because `format_num` had BOTH failure modes and only one of
    them was visible on the sheet: `"{:.4f}".format(None)` RAISES, and this is a per-page loop
    whose page guard tests only for an EMPTY response, so a present-but-null cell inside a
    healthy 200 cost the page.
    """
    for absent in (None, float('nan'), np.nan, '', 'nan', [], {}, True, False):
        assert pb._xl_num(absent) == 'N/A', repr(absent)
    assert pb._xl_num(0.0) == '0.0000'
    assert pb._xl_num(-1.5) == '-1.5000'
    assert pb._xl_num(3) == '3.0000'


def test_a_column_of_NULLS_cannot_kill_the_page(tmp_path, monkeypatch):
    """Removing `.fillna(0)` re-opens a page-kill path, so it is closed with the same change.

    pandas builds an OBJECT column of `None` from a response whose field is null on every row,
    and `col * 100` raises TypeError on it.  `.fillna(0)` was masking that on the dividend
    column -- the second reason it was there -- so the fix has to carry
    `pd.to_numeric(..., errors='coerce')` or it trades a false zero for a lost page.
    """
    km_all_null = [{'peRatio': None, 'ptbRatio': None, 'currentRatio': None,
                    'earningsYield': None, 'dividendYield': None,
                    'grahamNumber': None, 'freeCashFlowPerShare': None}] * 4
    grid = _write_xlsx(tmp_path, monkeypatch, overrides={'key-metrics': km_all_null})
    header = _table_header(grid)
    for col in (pb.XL_EARNYIELD_COL, pb.XL_PTB_COL, pb.XL_CURRATIO_COL, pb.XL_DIVYIELD_COL):
        assert grid[1][header.index(col)] == 'N/A', col


# =========================================================================== #
#  4.  THE SECTOR P/E WAS FETCHED FOR THE WRONG DATE                           #
# =========================================================================== #

def test_the_sector_PE_is_requested_for_the_date_the_code_MEANT(tmp_path, monkeypatch):
    """*** `strftime('%Y-%d-%m')` -- day and month swapped. ***

    `lastlast_5th` is always the 5th of some month, so the emitted string was always
    `YYYY-05-<real month number>` -- a real, parseable, in-range date every time, which is
    exactly why it survived: the endpoint answered, so nothing ever reported a failure.  Every
    run all year asked for a date in MAY, and the `Sector Average PE-ratio` cell on the CEO's
    pages was a May reading presented as a recent one.

    ASSERTED ON THE URL THE STAGE ACTUALLY BUILDS, not on the format string: a source scan
    would pass on any string containing the right tokens in any order, which is the mutation
    that got past the P/E's own earlier check.
    """
    import datetime as _dt
    seen = []
    _write_xlsx(tmp_path, monkeypatch, recorder=seen)
    urls = [u for u in seen if 'sector_price_earning_ratio' in u]
    assert urls, 'the sector-PE call was not made at all'
    asked = urls[0].split('date=')[1].split('&')[0]
    parsed = _dt.datetime.strptime(asked, '%Y-%m-%d').date()

    #  what the code MEANS: the 5th of the month five weeks back
    almago = _dt.date.today() - _dt.timedelta(weeks=5)
    want = (almago.replace(day=5) if almago.day >= 5
            else (almago - _dt.timedelta(days=almago.day)).replace(day=5))
    assert parsed == want, 'asked for %s, meant %s' % (parsed, want)
    #  the swap is what made the DAY always 5 -- name the symptom, not just the value
    assert parsed.day == 5
    assert asked == want.strftime('%Y-%m-%d')


# =========================================================================== #
#  5.  THE PEG COMMENT DESCRIBED A MECHANISM THAT WAS REMOVED THE SAME DAY     #
# =========================================================================== #

def test_the_PEG_seam_no_longer_describes_the_substitution_it_stopped_doing():
    """Stale prose on a live seam is a defect: the next reader trusts it.

    `substitute_peg_crossing` became a RETAINED NO-OP on 2026-09-03 (CEO: "refuse it, score as
    fail like every other undefined criterion"), and the comment at its call site still
    described the pool-median substitution as current -- including the audit-H-1 argument for
    WHY the call sits where it does, which rests on it being a cross-sectional baseline it no
    longer is.

    A SOURCE-TEXT TEST, AND THAT IS THE RIGHT INSTRUMENT HERE AND ONLY HERE: the defect IS the
    text.  It is a weaker instrument than the driven tests above and is not asked to carry
    anything else -- the BEHAVIOUR of the no-op is pinned by `test_nan_policy` and
    `test_peg_crossing_refusal`.
    """
    src = inspect.getsource(pb.postBoWrapper)
    i = src.index('substitute_peg_crossing')
    block = src[max(0, i - 2600):i]
    assert 'RETAINED NO-OP' in block or 'IS GONE' in block, (
        'the call site still presents the substitution as something that happens'
    )
    assert '2026-09-03' in block, 'the ruling that removed it is not named'
    #  the specific false claim, which is what a reader would have acted on
    assert "takes the POOL's median growth rate instead" not in block
    #  and the justification that no longer applies must not be left standing unqualified
    if 'audit H-1' in block:
        assert 'no longer applies' in block or 'no longer a cross-sectional' in block.lower()


def test_the_call_and_its_position_are_still_pinned():
    """The no-op stays where it is on purpose, so the prose fix must not have moved it.

    `test_nan_policy` pins both that it appears exactly once in this file and that it precedes
    Stage-1 scoring; this asserts the same two facts from the other side, so a prose edit that
    quietly dropped the call fails here rather than in an unrelated file.
    """
    #  `inspect.getsource(pb)`, not a read of `REPO/postBo.py`: the module under test is
    #  whatever is IMPORTED, and reading the path instead is how a mutation run silently
    #  checks the fixed file while claiming to check the unfixed one.
    src = inspect.getsource(pb)
    assert src.count('substitute_peg_crossing') == 1
    assert src.index('substitute_peg_crossing') < src.index('cs.simpleScore_fromDict')
